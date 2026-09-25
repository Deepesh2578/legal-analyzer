import os 
import re
import logging
from pathlib import Path
from typing import Dict, Any, List, Optional

# PDF libraries
import PyPDF2
import pdfplumber
import docx

logger = logging.getLogger(__name__)

class DocumentParser:
    def __init__(self):
        self.supported_formats = ['.pdf', '.docx']

    def parse_document(self, file_path: str, file_content: bytes = None) -> Dict[str, Any]:
        file_ext = Path(file_path).suffix.lower()
        if file_ext not in self.supported_formats:
            raise ValueError(f"Unsupported file format: {file_ext}")

        result = {
            "raw_text": "",
            "processed_text": "",
            "page_count": 0,
            "word_count": 0,
            "extraction_method": "",
            "extra_data": {}
        }

        if file_ext == '.pdf':
            result = self._parse_pdf(file_path, file_content)
        elif file_ext == '.docx':
            result = self._parse_docx(file_path, file_content)

        # Clean and count
        if result["raw_text"]:
            cleaned = self._clean_text(result["raw_text"])
            result["processed_text"] = cleaned
            result["word_count"] = len(cleaned.split())

        return result

    def _parse_pdf(self, file_path: str, file_content: bytes = None) -> Dict[str, Any]:
        # If content is provided, write to temp file for parsing
        if file_content:
            import tempfile
            with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as tmp:
                tmp.write(file_content)
                tmp_path = tmp.name
        else:
            tmp_path = file_path

        try:
            # Try pdfplumber first
            with pdfplumber.open(tmp_path) as pdf:
                text = "\n".join([page.extract_text() or "" for page in pdf.pages])
                metadata = {"title": pdf.metadata.get("Title", "")}
                return {
                    "raw_text": text,
                    "page_count": len(pdf.pages),
                    "extraction_method": "pdfplumber",
                    "extra_data": metadata
                }
        except Exception as e:
            logger.warning(f"pdfplumber failed: {e}. Falling back to PyPDF2.")
            try:
                with open(tmp_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    text = "\n".join([page.extract_text() or "" for page in reader.pages])
                    metadata = {"title": reader.metadata.get("/Title", "")}
                    return {
                        "raw_text": text,
                        "page_count": len(reader.pages),
                        "extraction_method": "pypdf2",
                        "extra_data": metadata
                    }
            except Exception as e2:
                logger.error(f"Both pdfplumber and PyPDF2 failed: {e2}")
                raise ValueError("Could not extract text from PDF")
        finally:
            if file_content and os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def _parse_docx(self, file_path: str, file_content: bytes = None) -> Dict[str, Any]:
        if file_content:
            import io
            doc = docx.Document(io.BytesIO(file_content))
        else:
            doc = docx.Document(file_path)

        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        text = "\n\n".join(paragraphs)
        return {
            "raw_text": text,
            "page_count": 1,
            "extraction_method": "docx",
            "extra_data": {"paragraph_count": len(paragraphs)}
        }

    def _clean_text(self, text: str) -> str:
        text = re.sub(r'\s+', ' ', text)
        text = re.sub(r'\n\s*\n', '\n\n', text)
        return text.strip()
